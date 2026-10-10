# -*- coding: utf-8 -*-
"""Excel 版の生成（Meiryo UI・モノトーン）

  - 本文は content.py、見た目の部品は xl_style.py（以前の資料と同じ部品）
  - 'seq' ブロック: 時刻 × 登場人物（JVM・カーネル・EFS・他タスク）の表で「動作イメージ」を描く
"""
from openpyxl.styles import Alignment
from openpyxl.worksheet.hyperlink import Hyperlink

import xl_style as X
from xl_style import F, FILL, WHITE, G03, G12, G55, G70, G80, INK, INK2
import content as C


# ---------------------------------------------------------------------------
# 動作イメージ（シーケンス表）
#   lanes: 列の見出し（先頭は「時刻」）
#   rows : (種類, [セル...])  種類 '' = 普通 / 'k' = 要所（薄灰・太字） / 'x' = 失敗（濃灰・白字）
# ---------------------------------------------------------------------------
def seq(ws, r, lanes, rows, widths=None, size=9.5):
    n = len(lanes)
    spans = X.alloc(widths if widths else [2] + [5] * (n - 1))
    starts = [X.C0 + sum(spans[:i]) for i in range(n)]
    for i, h in enumerate(lanes):
        X._cell(ws, r, starts[i], spans[i], h, F(10, True, WHITE), FILL(G55), X.B_BOX, X.AL_CTR)
        X.reg_h(ws, r, starts[i], spans[i], h, 10, 1.5, pad=8, minimum=30)
    ws.row_dimensions[r].height = 30
    r += 1
    for j, (kind, cells) in enumerate(rows):
        base = FILL(G03) if j % 2 == 1 else FILL(WHITE)
        for i in range(n):
            v = cells[i] if i < len(cells) else ""
            fill, font = base, F(size, i == 0, INK if i else INK2)
            if v and kind == 'x' and i > 0:
                fill, font = FILL(G70), F(size, True, WHITE)
            elif v and kind == 'k' and i > 0:
                fill, font = FILL(G12), F(size, True, INK)
            al = Alignment(horizontal="center" if i == 0 else "left", vertical="center",
                           wrap_text=True, indent=0 if i == 0 else 1)
            X._cell(ws, r, starts[i], spans[i], v or None, font, fill, X.B_BOX, al)
            if v:
                X.reg_h(ws, r, starts[i], spans[i], v, size, 1.5 if i == 0 else 3.0, pad=6, minimum=20)
        ws.row_dimensions[r].height = 20
        r += 1
    return r


# ---------------------------------------------------------------------------
# 汎用ブロックの描画
# ---------------------------------------------------------------------------
def render_blocks(ws, r, blocks):
    for b in blocks:
        kind = b[0]
        if kind == 'sec':
            r = X.gap(ws, r, 8)
            r = X.sec(ws, r, b[1])
            r = X.gap(ws, r, 4)
        elif kind == 'sub':
            r = X.gap(ws, r, 2)
            r = X.sub(ws, r, b[1])
        elif kind == 'p':
            r = X.body(ws, r, b[1])
        elif kind == 'ul':
            for item in b[1]:
                r = X.body(ws, r, "・" + item, indent=1)
            r = X.gap(ws, r, 4)
        elif kind == 'table':
            _, headers, rows, widths, align = b[:5]
            opts = b[5] if len(b) > 5 else {}
            # 数字だけの文字列（順位・段など）は数値で書く（Excel の「文字列として保存された数値」の緑の三角を出さない）
            rows = [[int(v) if isinstance(v, str) and v.isdigit() and len(v) < 6 else v for v in row] for row in rows]
            r = X.table(ws, r, headers, rows, widths=widths, align=align, check=opts.get('check', False))
            r = X.gap(ws, r, 6)
        elif kind == 'code':
            _, title, lines = b
            lines = [int(v) if v.isdigit() and len(v) < 6 else v for v in lines]   # 同上（緑の三角を出さない）
            r = X.code(ws, r, lines, title=title)
            r = X.gap(ws, r, 6)
        elif kind == 'callout':
            r = X.callout(ws, r, "■ " + b[1], b[2])
            r = X.gap(ws, r, 6)
        elif kind == 'flow':
            r = X.flow(ws, r, b[1])
            r = X.gap(ws, r, 6)
        elif kind == 'two':
            _, lt, ll, rt, rl = b
            r = X.twocol(ws, r, lt, ll, rt, rl, lw=1, rw=1)
            r = X.gap(ws, r, 6)
        elif kind == 'boxes':
            for row in b[1]:
                small = all(p[0] in ('note', 'space', 'arrow') for p in row)
                r = X.boxes(ws, r, row, h=22 if small else 46, size=9.5)
            r = X.gap(ws, r, 6)
        elif kind == 'seq':
            _, lanes, rows = b[:3]
            widths = b[3] if len(b) > 3 else None
            r = seq(ws, r, lanes, rows, widths)
            r = X.gap(ws, r, 6)
        else:
            raise ValueError(kind)
    return r


def link_urls(ws):
    """http(s):// で始まるセルをハイパーリンクにする（参考資料など）"""
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith(("http://", "https://")):
                c.hyperlink = v
                f = F(c.font.sz or 9.5, False, INK)
                f.underline = 'single'
                c.font = f


def new_sheet(wb, ch):
    ws = wb.create_sheet(ch['name'])
    X.sheet_setup(ws)
    r = X.head(ws, 2, ch['no'], ch['title'], ch.get('lead'))
    return ws, r


# ---------------------------------------------------------------------------
# 00 表紙
# ---------------------------------------------------------------------------
def cover(wb):
    ws = wb.active
    ws.title = "00_表紙"
    X.sheet_setup(ws, tab=G80)
    r = X.head(ws, 2, "00", C.TITLE, C.SUBTITLE + "\n" + C.AS_OF)
    r = X.sec(ws, r, "ご質問への答え（要約）")
    r = X.gap(ws, r, 4)
    r = X.table(ws, r, ["質問", "答え"], [list(x) for x in C.COVER_SUMMARY], widths=[5, 19])
    r = X.gap(ws, r, 8)
    r = X.sec(ws, r, "目次（シート名をクリックすると移動）")
    r = X.gap(ws, r, 4)
    first = r + 1
    r = X.table(ws, r, ["シート", "内容"], [list(x) for x in C.TOC], widths=[6, 18])
    for i, (name, _) in enumerate(C.TOC):
        c = ws.cell(row=first + i, column=X.C0)
        c.hyperlink = Hyperlink(ref=c.coordinate, location="'%s'!A1" % name, display=name)
        f = F(9.5, True, INK)
        f.underline = 'single'
        c.font = f
    r = X.gap(ws, r, 8)
    r = X.sec(ws, r, "凡例")
    r = X.gap(ws, r, 4)
    r = X.table(ws, r, ["表記", "意味"], [list(x) for x in C.LEGEND], widths=[6, 18])
    r = X.gap(ws, r, 8)
    X.body(ws, r, "作成: %s  ／  フォント: Meiryo UI  ／  配色: モノトーン  ／  本文は Markdown 版と同じ" % C.CREATED,
           size=9, color=INK2)


def build(path):
    wb = X.new_workbook()
    wb.properties.title = C.TITLE
    wb.properties.subject = C.SUBTITLE
    wb.properties.creator = "Claude Code"
    cover(wb)
    for ch in C.CHAPTERS:
        ws, r = new_sheet(wb, ch)
        render_blocks(ws, r, ch['blocks'])
        link_urls(ws)
        ws.sheet_view.selection[0].activeCell = "A1"
        ws.sheet_view.selection[0].sqref = "A1"
    n = X.finalize_heights()
    wb.active = 0
    wb.save(path)
    return n
